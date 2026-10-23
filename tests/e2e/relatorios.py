"""Fase 3.1: central de relatórios + exportação real (CSV, Excel, PDF). Abre os arquivos gerados e confere o conteúdo."""
import sys, pathlib, re, subprocess, csv, io
from playwright.sync_api import sync_playwright
import openpyxl
URL = pathlib.Path('dist/index.html').resolve().as_uri() + '?modo=local'  # sem nuvem nos testes de tela
SHOTS = pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else None
TMP = pathlib.Path('/tmp/claude-0/-home-claude/2892429b-2f99-5af8-bd4a-dce4628ac0a3/scratchpad/exports'); TMP.mkdir(parents=True, exist_ok=True)
res = []
def ok(n, c): res.append(('PASSOU' if c else 'FALHOU', n)); return c
IDS = ['os', 'fazenda', 'aeronave', 'operador', 'planejado', 'clima', 'paradas', 'evolucao', 'tecnico']

def rodar(vp, rot, mobile=False):
    with sync_playwright() as p:
        b = p.chromium.launch()
        ctx = b.new_context(viewport=vp, is_mobile=mobile, has_touch=mobile, accept_downloads=True)
        pg = ctx.new_page(); erros = []; pg.on('pageerror', lambda e: erros.append(str(e)))
        rede = []; pg.on('request', lambda r: rede.append(r.url) if not r.url.startswith(('file:', 'blob:', 'data:')) and 'fonts.g' not in r.url else None)
        pg.goto(URL + '#/relatorios'); pg.wait_for_selector('.rel-lista')
        ok(f'[{rot}] central lista Sequência, Diário e os 9 relatórios (nenhum "planejado")', pg.locator('.rel-lista a.rel-item').count() == 11 and pg.locator('.rel-planejado').count() == 0)
        for rid in IDS:
            pg.goto(URL + f'#/relatorios/ver?id={rid}'); pg.wait_for_selector('.rel-doc')
            cab = pg.locator('.rel-cab').inner_text()
            ok(f'[{rot}] {rid}: tela com título, filtros, gerado em e responsável', 'Filtros:' in cab and 'Gerado em' in cab and 'Responsável' in cab and (rid == 'tecnico' or 'Período' in cab))
        # Responsável editável e levado aos arquivos
        pg.goto(URL + '#/relatorios/ver?id=os'); pg.wait_for_selector('.rel-doc')
        pg.fill('[data-k="resp"]', 'Eng. Responsável Teste'); pg.locator('[data-k="resp"]').dispatch_event('change'); pg.wait_for_timeout(150)
        ok(f'[{rot}] responsável aparece no relatório', 'Eng. Responsável Teste' in pg.locator('.rel-cab').inner_text())
        tot_tela = pg.locator('.rel-t tfoot td').nth(6).inner_text()
        arqs = {}
        for fmt in ['csv', 'xlsx', 'pdf']:
            with pg.expect_download(timeout=60000) as dl: pg.click(f'[data-exp="{fmt}"]')
            caminho = TMP / f'{rot}-{dl.value.suggested_filename}'; dl.value.save_as(str(caminho)); arqs[fmt] = caminho
        ok(f'[{rot}] 3 arquivos com nome do relatório e período', all(re.match(r'.*resumo-por-ordem-de-servico_\d{4}-\d{2}-\d{2}_a_\d{4}-\d{2}-\d{2}\.' + f, str(arqs[f])) for f in arqs))
        txt = arqs['csv'].read_text(encoding='utf-8-sig')
        linhas = list(csv.reader(io.StringIO(txt), delimiter=';'))
        ok(f'[{rot}] CSV: cabeçalho (período, filtros, gerado, responsável) e total igual à tela', linhas[0][0] == 'Resumo por ordem de serviço' and linhas[1][0].startswith('Período:') and 'Eng. Responsável Teste' in linhas[3][0] and any(l and l[0] == 'Total' and l[6] == tot_tela.replace('.', '') for l in linhas))  # CSV sem separador de milhar (Excel soma direto)
        wb = openpyxl.load_workbook(arqs['xlsx']); ws = wb.active
        vals = [[c.value for c in r] for r in ws.iter_rows()]
        tot = [r for r in vals if r and r[0] == 'Total'][0]
        ok(f'[{rot}] Excel abre, tem cabeçalho e total NUMÉRICO igual à tela', vals[0][0] == 'Resumo por ordem de serviço' and isinstance(tot[6], (int, float)) and f'{tot[6]:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.') == tot_tela and any('Eng. Responsável Teste' in str(r[0]) for r in vals[:5]))
        ok(f'[{rot}] Excel: datas como data e % como fração', any(isinstance(r[12], __import__('datetime').datetime) for r in vals if len(r) > 12) and ws.cell(row=[i for i, r in enumerate(vals, 1) if r and r[0] == 'Total'][0], column=11).number_format == '0.0%')
        pdftxt = subprocess.run(['pdftotext', '-layout', str(arqs['pdf']), '-'], capture_output=True, text=True).stdout
        info = subprocess.run(['pdfinfo', str(arqs['pdf'])], capture_output=True, text=True).stdout
        ok(f'[{rot}] PDF válido com título, período, responsável, total e paginação', 'Resumo por ordem de serviço' in pdftxt and 'Período:' in pdftxt and 'Eng. Responsável Teste' in pdftxt and tot_tela in pdftxt and re.search(r'Página 1 de \d+', pdftxt) and 'Pages:' in info)
        # Relatório técnico em PDF (retrato) com produtos e clima
        pg.goto(URL + '#/relatorios/ver?id=tecnico'); pg.wait_for_selector('.rel-doc')
        ok(f'[{rot}] técnico: escolhe a operação e mostra identificação, calda, clima, paradas', pg.locator('[data-k="op"] option').count() > 5 and all(x in pg.locator('.rel-doc').inner_text() for x in ['Identificação', 'Calda', 'Condições climáticas', 'Paradas', 'Observações']))
        with pg.expect_download(timeout=60000) as dl: pg.click('[data-exp="pdf"]')
        c2 = TMP / f'{rot}-tecnico.pdf'; dl.value.save_as(str(c2))
        t2 = subprocess.run(['pdftotext', '-layout', str(c2), '-'], capture_output=True, text=True).stdout
        info2 = subprocess.run(['pdfinfo', str(c2)], capture_output=True, text=True).stdout
        ok(f'[{rot}] PDF técnico em retrato com dose e avaliação climática', 'Relatório técnico da aplicação' in t2 and 'Dose planejada' in t2 and ('Dentro' in t2 or 'Fora' in t2) and '595.28 x 841.89' in info2)
        # Filtro muda os números (fazenda)
        pg.goto(URL + '#/relatorios/ver?id=planejado'); pg.wait_for_selector('.rel-doc')
        antes = pg.locator('.rel-kpis').inner_text()
        sel = pg.locator('#rel-filtros [data-f="fazendaId"]')
        if sel.count():
            if mobile and pg.locator('#rel-filtros .filtros-toggle, #rel-filtros [data-acao="filtros"]').count(): pg.locator('#rel-filtros .filtros-toggle, #rel-filtros [data-acao="filtros"]').first.click()
            v = sel.locator('option').nth(1).get_attribute('value'); sel.select_option(v); pg.wait_for_timeout(200)
            ok(f'[{rot}] filtro de fazenda altera o relatório e aparece no cabeçalho', pg.locator('.rel-kpis').inner_text() != antes and 'Fazenda:' in pg.locator('.rel-cab').inner_text())
        if SHOTS: pg.screenshot(path=str(SHOTS / f'{rot}-relatorio.png'), full_page=True)
        # Sequência e Diário também exportam
        for rota, titulo in [('relatorios/sequencia', 'Sequência de aplicação'), ('relatorios/diario', 'Diário de aplicações')]:
            pg.goto(URL + '#/' + rota); pg.wait_for_selector('[data-exp="pdf"]')
            for fmt in ['xlsx', 'pdf']:
                with pg.expect_download(timeout=60000) as dl: pg.locator(f'[data-exp="{fmt}"]').first.click()
                c3 = TMP / f'{rot}-{dl.value.suggested_filename}'; dl.value.save_as(str(c3))
                if fmt == 'pdf':
                    t3 = subprocess.run(['pdftotext', '-layout', str(c3), '-'], capture_output=True, text=True).stdout
                    ok(f'[{rot}] {titulo}: PDF com título, responsável e tabela', titulo in t3 and 'Eng. Responsável Teste' in t3 and ('Total' in t3))
                else:
                    wb3 = openpyxl.load_workbook(c3); ok(f'[{rot}] {titulo}: Excel abre', str(wb3.active.cell(1, 1).value).startswith(titulo))
        ok(f'[{rot}] exportação 100% local (nenhum acesso à rede)', not rede)
        ok(f'[{rot}] sem erros de JavaScript', not erros)
        if erros: print(erros)
        if rede: print(rede)
        b.close()

rodar({'width': 1440, 'height': 900}, 'desktop')
rodar({'width': 390, 'height': 844}, 'iphone', True)
for s, n in res: print(f'{s}  {n}')
print(f"\n{sum(1 for s,_ in res if s=='PASSOU')}/{len(res)} verificações passaram")
sys.exit(0 if all(s == 'PASSOU' for s, _ in res) else 1)
