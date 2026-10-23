"""Fase 3.2: análise de produtividade (A × B), quebra por dimensão, exportação."""
import sys, pathlib, subprocess
from playwright.sync_api import sync_playwright
import openpyxl
URL = pathlib.Path('dist/index.html').resolve().as_uri()
SHOTS = pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else None
TMP = pathlib.Path('/tmp/claude-0/-home-claude/2892429b-2f99-5af8-bd4a-dce4628ac0a3/scratchpad/exports'); TMP.mkdir(parents=True, exist_ok=True)
res = []
def ok(n, c): res.append(('PASSOU' if c else 'FALHOU', n)); return c
def rodar(vp, rot, mobile=False):
    with sync_playwright() as p:
        b = p.chromium.launch(); ctx = b.new_context(viewport=vp, is_mobile=mobile, has_touch=mobile, accept_downloads=True)
        pg = ctx.new_page(); erros = []; pg.on('pageerror', lambda e: erros.append(str(e)))
        pg.goto(URL + '#/analise'); pg.wait_for_selector('.an-t')
        ok(f'[{rot}] período A cobrindo tudo: avisa que B está vazio', 'sem registros' in pg.locator('#an-corpo').inner_text())
        pg.select_option('[data-k="rapido"]', '14'); pg.wait_for_selector('.an-t'); pg.wait_for_timeout(300)
        ok(f'[{rot}] "últimos 14 dias" compara com os 14 anteriores (B com dados)', 'sem registros' not in pg.locator('#an-corpo').inner_text())
        t = pg.locator('.an-t').first.inner_text().lower()
        ok(f'[{rot}] tabela A × B com ha/h, eficiência, parado e variação', all(x in t for x in ['ha/h efetivo', 'eficiência', '% do tempo parado', 'variação']) and ('▲' in t or '▼' in t))
        ok(f'[{rot}] gráficos desenhados', pg.locator('#g-an-prod svg').count() == 1 and pg.locator('#g-an-area svg').count() == 1)
        for d, rotulo in [('pilotoId', 'Piloto'), ('turnoId', 'Turno'), ('fazenda', 'Fazenda')]:
            pg.select_option('[data-k="dim"]', d); pg.wait_for_timeout(150)
            ok(f'[{rot}] quebra por {rotulo}', rotulo.lower() in pg.locator('.an-t').nth(1).inner_text().lower() and pg.locator('.an-t').nth(1).locator('tbody tr').count() > 1)
        per0 = pg.locator('.an-per').inner_text()
        pg.select_option('[data-k="modo"]', 'semana'); pg.wait_for_timeout(150)
        ok(f'[{rot}] trocar o período B muda as datas', pg.locator('.an-per').inner_text() != per0)
        pg.select_option('[data-k="modo"]', 'livre'); pg.wait_for_timeout(100)
        ok(f'[{rot}] período livre mostra os campos de data', pg.locator('[data-k="de"]').is_visible())
        pg.select_option('[data-k="modo"]', 'anterior'); pg.wait_for_timeout(100)
        if SHOTS: pg.screenshot(path=str(SHOTS / f'{rot}-analise.png'), full_page=True)
        with pg.expect_download(timeout=60000) as dl: pg.click('[data-exp="xlsx"]')
        x = TMP / f'{rot}-analise.xlsx'; dl.value.save_as(str(x))
        ws = openpyxl.load_workbook(x).active
        ok(f'[{rot}] Excel do comparativo com períodos A e B', str(ws.cell(1, 1).value).startswith('Análise de produtividade') and any('Período B' in str(r[0].value) for r in ws.iter_rows(max_row=5)))
        with pg.expect_download(timeout=60000) as dl: pg.click('[data-exp="pdf"]')
        y = TMP / f'{rot}-analise.pdf'; dl.value.save_as(str(y))
        ok(f'[{rot}] PDF do comparativo', 'Análise de produtividade' in subprocess.run(['pdftotext', str(y), '-'], capture_output=True, text=True).stdout)
        pg.goto(URL + '#/planejamento'); pg.wait_for_selector('.plan-c')
        t = pg.locator('main, #app').first.inner_text().lower()
        menu = pg.locator('a[href="#/planejamento"]').first.inner_text().lower() if pg.locator('a[href="#/planejamento"]').count() else ''
        ok(f'[{rot}] Planejamento: metas vigentes e um cartão por campanha com prazo e necessário (sem selo "planejado")', pg.locator('.plan-c').count() >= 2 and 'meta diária' in t and 'necessário para o prazo' in t and 'planejado' not in menu)
        with pg.expect_download(timeout=60000) as dl: pg.click('[data-exp="pdf"]')
        z = TMP / f'{rot}-plan.pdf'; dl.value.save_as(str(z))
        ok(f'[{rot}] PDF do planejamento', 'Planejamento e metas' in subprocess.run(['pdftotext', str(z), '-'], capture_output=True, text=True).stdout)
        ok(f'[{rot}] sem erros de JavaScript', not erros)
        if erros: print(erros)
        b.close()
rodar({'width': 1440, 'height': 900}, 'desktop')
rodar({'width': 390, 'height': 844}, 'iphone', True)
for s, n in res: print(f'{s}  {n}')
print(f"\n{sum(1 for s,_ in res if s=='PASSOU')}/{len(res)} verificações passaram")
sys.exit(0 if all(s == 'PASSOU' for s, _ in res) else 1)
